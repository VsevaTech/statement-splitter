"""Exact monetary arithmetic: Decimal end-to-end (parse → model → SQLite → summary → export)."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app import categories as c
from app import services
from app.db import Statement, Transaction
from app.export import build_summary, export_workbook, safe_text
from app.integrity import currency_totals
from app.money import canonical, dsum, excel_number_format, format_money, to_decimal
from app.parsing import parse_amount
from app.schemas import NormalizedTransaction
from tests.conftest import csv_bytes


def _tx(i, amount, currency="ILS", category=c.OFFICE, desc="SHOP"):
    amount = Decimal(amount)
    return NormalizedTransaction(
        row_index=i,
        date=None,
        description=desc,
        normalized_merchant=desc,
        amount=amount,
        currency=currency,
        direction="credit" if amount > 0 else "debit",
        category=category,
        category_source="builtin",
    )


def _import(session, body: str, name="s.csv") -> Statement:
    statement, _pending, mapping = services.import_statement(session, name, csv_bytes(body))
    assert statement is not None, mapping
    return statement


def test_point_one_plus_point_two_is_exact():
    assert 0.1 + 0.2 != 0.3  # the float problem this app avoids
    assert parse_amount("0.1") + parse_amount("0.2") == Decimal("0.3")
    assert dsum([parse_amount("0.10"), parse_amount("0.20"), parse_amount("0.30")]) == Decimal("0.60")


def test_critical_ten_twenty_thirty_cents_through_the_whole_pipeline(session):
    st = _import(
        session,
        "Date,Description,Amount,Currency\n"
        "2024-09-01,FX COMMISSION 1,0.10,ILS\n2024-09-02,FX COMMISSION 2,0.20,ILS\n2024-09-03,FX COMMISSION 3,0.30,ILS\n",
    )
    txs = services.list_transactions(session, st.id)
    total = dsum(t.amount for t in txs)
    assert total == Decimal("0.60")
    assert canonical(total) == "0.60"  # not 0.6000000000000001
    [totals] = currency_totals(txs)
    assert canonical(totals.net) == "0.60"


def test_large_monetary_value_keeps_every_cent(session):
    st = _import(
        session,
        'Date,Description,Amount,Currency\n2024-09-01,BIG DEAL,"98,765,432,109,876.54",ILS\n'
        "2024-09-02,SMALL,0.01,ILS\n",
    )
    total = dsum(t.amount for t in services.list_transactions(session, st.id))
    assert total == Decimal("98765432109876.55")
    assert parse_amount("98 765 432 109 876,54") == Decimal("98765432109876.54")


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("1234.56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("1 234.56", "1234.56"),
        ("1.234,56", "1234.56"),
        ("1234,56", "1234.56"),
        ("-125.40", "-125.40"),
        ("(125.40)", "-125.40"),
        ("125.40-", "-125.40"),
        ("125.40 DR", "-125.40"),
        ("125.40 CR", "125.40"),
        ("1'234.50", "1234.50"),
        ("12.345", "12.345"),  # 3-decimal currencies are not rounded to cents
        ("₪ 45.90", "45.90"),
        ("USD 12.00", "12.00"),
    ],
)
def test_amount_formats_parse_to_exact_decimal(raw, expected):
    value = parse_amount(raw)
    assert isinstance(value, Decimal)
    assert value == Decimal(expected)
    assert canonical(value) == expected


def test_negative_amount_and_decimal_comma_and_thousands_separator():
    assert parse_amount("-1.234,50") == Decimal("-1234.50")
    assert parse_amount("-45,90") == Decimal("-45.90")
    assert parse_amount("1,234,567.89") == Decimal("1234567.89")
    assert parse_amount("1.234.567,89") == Decimal("1234567.89")


@pytest.mark.parametrize("raw", ["12.3O", "12abc", "1,2,3.4", "12,34.56", "1.234,5.6", "--", "12..5", "O.50"])
def test_malformed_amount_is_none_never_zero(raw):
    assert parse_amount(raw) is None


def test_spreadsheet_floats_become_their_shortest_decimal():
    assert to_decimal(0.1) == Decimal("0.1")
    assert parse_amount(120.5) == Decimal("120.5")
    assert parse_amount(46.8) == Decimal("46.8")


def test_many_small_transactions_do_not_accumulate_rounding_error(session):
    rows = "".join(f"2024-09-{1 + i % 28:02d},MICRO FEE {i},-0.01,ILS\n" for i in range(3000))
    rows += "".join(f"2024-09-{1 + i % 28:02d},MICRO REFUND {i},0.07,ILS\n" for i in range(1000))
    st = _import(session, "Date,Description,Amount,Currency\n" + rows)
    txs = services.list_transactions(session, st.id)
    assert len(txs) == 4000
    [t] = currency_totals(txs)
    assert t.expenses == Decimal("30.00")
    assert t.income == Decimal("70.00")
    assert t.net == Decimal("40.00")
    float_sum = sum(float(x.amount) for x in txs)
    assert float_sum != 40.0  # binary floats drift; the Decimal path does not


def test_category_aggregation_is_exact():
    txs = [_tx(i, "-0.10", category=c.BANK_FEES) for i in range(10)] + [_tx(10, "-33.33", category=c.OFFICE)] * 3
    rows = {(cat, cur): (n, total) for cat, n, total, cur in build_summary(txs)}
    assert rows[(c.BANK_FEES, "ILS")] == (10, Decimal("-1.00"))
    assert rows[(c.OFFICE, "ILS")] == (3, Decimal("-99.99"))
    assert all(isinstance(total, Decimal) for _c, _n, total, _cur in build_summary(txs))


def test_currency_aggregation_never_mixes_currencies():
    txs = [_tx(0, "-0.10", "ILS"), _tx(1, "-0.20", "ILS"), _tx(2, "-0.30", "USD"), _tx(3, "5.05", "EUR")]
    by_cur = {t.currency: t for t in currency_totals(txs)}
    assert set(by_cur) == {"ILS", "USD", "EUR"}
    assert by_cur["ILS"].net == Decimal("-0.30")
    assert by_cur["USD"].net == Decimal("-0.30")
    assert by_cur["EUR"].income == Decimal("5.05") and by_cur["EUR"].expenses == Decimal("0.00")


def test_database_round_trip_is_exact_and_stored_as_text(session):
    st = _import(
        session,
        "Date,Description,Amount,Currency\n2024-09-01,A,0.10,ILS\n2024-09-02,B,-1234567.891,KWD\n"
        "2024-09-03,C,-12.30,ILS\n",
    )
    session.expire_all()
    amounts = [t.amount for t in services.list_transactions(session, st.id)]
    assert amounts == [Decimal("0.10"), Decimal("-1234567.891"), Decimal("-12.30")]
    assert all(isinstance(a, Decimal) for a in amounts)
    raw = session.execute(text("SELECT amount, typeof(amount) FROM transactions ORDER BY row_index")).all()
    assert raw == [("0.10", "text"), ("-1234567.891", "text"), ("-12.30", "text")]


def test_database_rejects_binary_float_money(session):
    st = Statement(id="00000000-0000-0000-0000-000000000001", filename="x.csv")
    session.add(st)
    session.add(
        Transaction(
            statement_id=st.id,
            row_index=0,
            description="X",
            normalized_merchant="X",
            amount=0.1,
            currency="ILS",
            direction="credit",
            category_source="unknown",
        )
    )
    with pytest.raises(Exception, match="binary float"):
        session.commit()
    session.rollback()


def test_xlsx_export_writes_exact_numeric_cells(session):
    st = _import(
        session,
        "Date,Description,Amount,Currency\n2024-09-01,FX COMMISSION 1,-0.10,ILS\n"
        "2024-09-02,FX COMMISSION 2,-0.20,ILS\n2024-09-03,FX COMMISSION 3,-0.30,ILS\n"
        "2024-09-04,CLIENT,12.30,USD\n",
    )
    data = export_workbook(services.list_transactions(session, st.id), services.integrity_report(session, st))
    wb = load_workbook(io.BytesIO(data))
    amounts = [r[3].value for r in wb["Transactions"].iter_rows(min_row=2)]
    assert all(isinstance(v, (int, float)) for v in amounts)  # numeric cells, not text
    assert amounts == [-0.1, -0.2, -0.3, 12.3]
    assert "Decimal(" not in data.decode("latin-1")
    summary = {(r[0].value, r[3].value): r[2].value for r in wb["Summary"].iter_rows(min_row=2) if r[3].value}
    assert Decimal(str(summary[(c.BANK_FEES, "ILS")])) == Decimal("-0.60")
    assert wb["Transactions"]["D2"].number_format == "#,##0.00"


def test_xlsx_three_decimal_currency_keeps_its_scale():
    assert excel_number_format([Decimal("1.234"), Decimal("5.00")]) == "#,##0.000"
    assert excel_number_format([Decimal("5")]) == "#,##0.00"


def test_formula_injection_is_neutralized_in_export():
    evil = _tx(0, "-1.00", desc='=HYPERLINK("http://evil","x")')
    wb = load_workbook(io.BytesIO(export_workbook([evil])))
    cell = wb["Transactions"]["B2"]
    assert cell.data_type == "s" and cell.value.startswith("'=")
    assert safe_text("+1") == "'+1" and safe_text("@SUM(A1)") == "'@SUM(A1)" and safe_text("WOLT") == "WOLT"


def test_money_display_format():
    assert format_money(Decimal("12450.3")) == "12,450.30"
    assert format_money(Decimal("-12.40"), signed=True) == "-12.40"
    assert format_money(Decimal("6378.63"), signed=True) == "+6,378.63"
    assert format_money(Decimal("-0.00")) == "0.00"
    assert format_money(Decimal("1.234")) == "1.234"


def test_json_values_are_decimal_strings(client, session):
    st = _import(session, "Date,Description,Amount,Currency\n2024-09-01,A,0.10,ILS\n2024-09-02,B,0.20,ILS\n")
    data = client.get(f"/api/statements/{st.id}").json()
    [totals] = data["currency_totals"]
    assert totals["net_movement"] == "0.30" and isinstance(totals["net_movement"], str)


def test_existing_rows_query_still_works(session):
    _import(session, "Date,Description,Amount,Currency\n2024-09-01,UBER *TRIP 1,-30.00,ILS\n")
    tx = session.scalar(select(Transaction))
    assert tx.amount == Decimal("-30.00") and tx.category == c.TRANSPORT

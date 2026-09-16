from __future__ import annotations

import io

from openpyxl import load_workbook

from app import categories as c
from app.export import SUMMARY_HEADERS, TRANSACTION_HEADERS, build_summary, export_workbook
from app.schemas import NormalizedTransaction


def _tx(i, desc, amount, currency, category, source="builtin"):
    return NormalizedTransaction(
        row_index=i,
        date=None,
        description=desc,
        normalized_merchant=desc.split(" ")[0],
        amount=amount,
        currency=currency,
        direction="credit" if amount > 0 else "debit",
        category=category,
        category_source=source,
        confidence=0.9,
    )


TXS = [
    _tx(0, "UBER 1", -30.0, "ILS", c.TRANSPORT),
    _tx(1, "UBER 2", -20.0, "ILS", c.TRANSPORT),
    _tx(2, "UBER 3", -10.0, "USD", c.TRANSPORT),
    _tx(3, "AWS", -100.0, "USD", c.CLOUD),
    _tx(4, "HETZNER", -40.0, "EUR", c.CLOUD),
    _tx(5, "CLIENT", 1000.0, "ILS", c.INCOME, "income"),
    _tx(6, "MYSTERY", -5.0, "ILS", None, "unknown"),
]


def test_multi_currency_summary_never_mixes_currencies():
    rows = build_summary(TXS)
    as_dict = {(cat, cur): (count, total) for cat, count, total, cur in rows}
    assert as_dict[(c.TRANSPORT, "ILS")] == (2, -50.0)
    assert as_dict[(c.TRANSPORT, "USD")] == (1, -10.0)
    assert as_dict[(c.CLOUD, "USD")] == (1, -100.0)
    assert as_dict[(c.CLOUD, "EUR")] == (1, -40.0)
    assert as_dict[(c.INCOME, "ILS")] == (1, 1000.0)
    assert as_dict[("Needs review", "ILS")] == (1, -5.0)
    assert len(rows) == 6


def test_export_xlsx_sheets_and_headers():
    wb = load_workbook(io.BytesIO(export_workbook(TXS)))
    assert wb.sheetnames == ["Transactions", "Summary"]
    ws = wb["Transactions"]
    assert tuple(cell.value for cell in ws[1]) == TRANSACTION_HEADERS
    assert ws.max_row == len(TXS) + 1
    row = [cell.value for cell in ws[2]]
    assert row[1:] == ["UBER 1", "UBER", -30.0, "ILS", "debit", c.TRANSPORT, "builtin"]
    assert [cell.value for cell in ws[8]][6] == "Needs review"

    ws2 = wb["Summary"]
    assert tuple(cell.value for cell in ws2[1]) == SUMMARY_HEADERS
    summary_rows = [tuple(cell.value for cell in r) for r in ws2.iter_rows(min_row=2)]
    currencies_per_category: dict[str, set[str]] = {}
    for cat, _count, _total, cur in summary_rows:
        currencies_per_category.setdefault(cat, set()).add(cur)
    assert currencies_per_category[c.TRANSPORT] == {"ILS", "USD"}
    assert currencies_per_category[c.CLOUD] == {"USD", "EUR"}

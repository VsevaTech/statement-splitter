"""Statement balance reconciliation: opening + net movement = closing, per currency."""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import load_workbook

from app import services
from app.columns import ColumnMapping
from app.integrity import BALANCED, DIFFERENCE, INCOMPLETE, balance_checks, currency_totals
from app.transform import analyze_rows
from tests.conftest import DEMO, csv_bytes


def _import(session, name=None, body=None):
    data = (DEMO / name).read_bytes() if name else csv_bytes(body)
    statement, _p, mapping = services.import_statement(session, name or "s.csv", data)
    assert statement is not None, mapping
    return statement


def _report(session, statement):
    session.expire_all()
    return services.integrity_report(session, session.merge(statement))


class _Bal:
    def __init__(self, currency, opening=None, closing=None):
        self.currency, self.opening, self.closing = currency, opening, closing
        self.opening_source = self.closing_source = "manual"
        self.opening_row = self.closing_row = None
        self.confirmed = True


class _Tx:
    def __init__(self, amount, currency="ILS"):
        self.amount, self.currency = Decimal(amount), currency


def test_net_movement_income_and_expenses():
    [t] = currency_totals([_Tx("8500.00"), _Tx("-2000.00"), _Tx("-121.37")])
    assert (t.income, t.expenses, t.net) == (Decimal("8500.00"), Decimal("2121.37"), Decimal("6378.63"))


def test_balanced_statement_math():
    totals = currency_totals([_Tx("8500.00"), _Tx("-2121.37")])
    [check] = balance_checks(totals, [_Bal("ILS", Decimal("12450.30"), Decimal("18828.93"))])
    assert check.expected_closing == Decimal("18828.93")
    assert check.difference == Decimal("0.00") and check.status == BALANCED


def test_unbalanced_statement_math():
    totals = currency_totals([_Tx("8500.00"), _Tx("-2121.37")])
    [check] = balance_checks(totals, [_Bal("ILS", Decimal("12450.30"), Decimal("18816.53"))])
    assert check.difference == Decimal("-12.40") and check.status == DIFFERENCE


def test_missing_opening_or_closing_balance_is_incomplete_not_guessed():
    totals = currency_totals([_Tx("10.00")])
    [no_open] = balance_checks(totals, [_Bal("ILS", None, Decimal("10.00"))])
    [no_close] = balance_checks(totals, [_Bal("ILS", Decimal("0.00"), None)])
    assert no_open.status == INCOMPLETE and no_open.missing == ["opening balance"] and no_open.difference is None
    assert no_close.status == INCOMPLETE and no_close.missing == ["closing balance"]
    assert no_close.expected_closing == Decimal("10.00")


def test_balanced_fixture(session, client):
    st = _import(session, "statement-balanced.csv")
    data = client.get(f"/api/statements/{st.id}").json()
    assert data["import_summary"] == {
        "source_rows": 5,
        "imported": 3,
        "ignored": 2,
        "needs_attention": 0,
        "not_imported": 0,
        "imported_with_warnings": 0,
    }
    assert data["balance_check"] == {
        "currency": "ILS",
        "opening": "1000.00",
        "income": "500.10",
        "expenses": "300.50",
        "net_movement": "199.60",
        "expected_closing": "1199.60",
        "actual_closing": "1199.60",
        "difference": "0.00",
        "status": "balanced",
        "opening_source": "detected",
        "closing_source": "detected",
        "confirmed": False,
    }
    page = client.get(f"/statements/{st.id}").text
    assert "Statement balance reconciled" in page and "Difference: 0.00 ILS" in page


def test_unbalanced_fixture_shows_known_difference(session, client):
    st = _import(session, "statement-unbalanced.csv")
    check = client.get(f"/api/statements/{st.id}").json()["balance_check"]
    assert (check["expected_closing"], check["actual_closing"], check["difference"], check["status"]) == (
        "1199.60",
        "1187.20",
        "-12.40",
        "difference",
    )
    page = client.get(f"/statements/{st.id}").text
    assert "Statement balance difference" in page and "-12.40 ILS" in page and "Needs review" in page
    assert "Accounting verified" not in page and "guaranteed" not in page


def test_demo_statement_is_balanced_in_ils_only(session):
    st = _import(session, "statement.csv")
    report = _report(session, st)
    [check] = report.balance_checks
    assert check.currency == "ILS" and check.status == BALANCED and check.difference == Decimal("0.00")
    assert report.currencies_without_balance == ["EUR", "USD"]


def test_multi_currency_balances_are_checked_separately(session):
    body = (
        "Date,Description,Amount,Currency\n"
        "2024-09-01,Opening balance,100.00,ILS\n2024-09-01,Opening balance,50.00,USD\n"
        "2024-09-02,WOLT IL,-10.00,ILS\n2024-09-03,AWS EMEA,-5.00,USD\n2024-09-04,CLIENT,20.00,USD\n"
        "2024-09-30,Closing balance,90.00,ILS\n2024-09-30,Closing balance,64.00,USD\n"
    )
    st = _import(session, body=body)
    checks = {c.currency: c for c in _report(session, st).balance_checks}
    assert checks["ILS"].status == BALANCED
    assert checks["USD"].net_movement == Decimal("15.00")
    assert checks["USD"].difference == Decimal("-1.00") and checks["USD"].status == DIFFERENCE


def test_ambiguous_balance_rows_are_not_used():
    cols = ["Date", "Description", "Amount", "Currency"]
    rows = [
        ["2024-09-01", "Opening balance", "100.00", "ILS"],
        ["2024-09-15", "Opening balance", "130.00", "ILS"],
        ["2024-09-02", "WOLT IL", "-10.00", "ILS"],
    ]
    m = ColumnMapping(date="Date", description="Description", amount="Amount", currency="Currency")
    a = analyze_rows(cols, rows, m)
    assert a.balances == [] and a.ambiguous_balances == ["opening ILS"]
    assert a.ignored == 2  # still accounted for


def test_manual_balance_entry(session, client):
    st = _import(session, body="Date,Description,Amount,Currency\n2024-09-02,WOLT IL,-45.90,ILS\n")
    assert "Balance check not run" in client.get(f"/statements/{st.id}").text
    resp = client.post(
        f"/statements/{st.id}/balances",
        data={"currency": "ILS", "opening": "1,000.00", "closing": "954.10"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    check = client.get(f"/api/statements/{st.id}").json()["balance_check"]
    assert check["status"] == "balanced" and check["opening_source"] == "manual" and check["confirmed"] is True


def test_manual_balance_rejects_garbage_instead_of_zero(session, client):
    st = _import(session, body="Date,Description,Amount,Currency\n2024-09-02,WOLT IL,-45.90,ILS\n")
    resp = client.post(
        f"/statements/{st.id}/balances",
        data={"currency": "ILS", "opening": "12.3O", "closing": ""},
        follow_redirects=False,
    )
    assert "balance_error=" in resp.headers["location"]
    assert client.get(f"/api/statements/{st.id}").json()["balance_checks"] == []


def test_confirm_detected_balances_keeps_detected_source(session, client):
    st = _import(session, "statement-balanced.csv")
    client.post(f"/statements/{st.id}/balances", data={"currency": "ILS", "opening": "1000.00", "closing": "1199.60"})
    check = client.get(f"/api/statements/{st.id}").json()["balance_check"]
    assert check["confirmed"] is True and check["opening_source"] == "detected" and check["status"] == "balanced"


def test_debit_credit_columns_direction_semantics(session):
    body = (
        "Date,Details,Debit,Credit,Currency\n2024-09-01,Opening balance,,,ILS\n"
        "2024-09-02,CLIENT PAYMENT,,500.10,ILS\n2024-09-03,OFFICE DEPOT 1,200.20,,ILS\n"
        "2024-09-04,WOLT IL,100.30,,ILS\n"
    )
    st = _import(session, body=body)
    [t] = currency_totals(services.list_transactions(session, st.id))
    assert (t.income, t.expenses, t.net) == (Decimal("500.10"), Decimal("300.50"), Decimal("199.60"))


def test_dr_cr_indicator_direction_semantics():
    cols = ["When", "Who", "Value", "DrCr"]
    rows = [["2024-09-01", "CLIENT", "1000.00", "CR"], ["2024-09-02", "WOLT IL", "45.90", "DR"]]
    m = ColumnMapping(date="When", description="Who", amount="Value", direction="DrCr")
    [t] = currency_totals(analyze_rows(cols, rows, m).transactions)
    assert t.net == Decimal("954.10") and t.expenses == Decimal("45.90")


def test_export_summary_contains_integrity_and_balance(session, client):
    st = _import(session, "statement-unbalanced.csv")
    wb = load_workbook(io.BytesIO(client.get(f"/statements/{st.id}/export").content))
    values = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2) if r[0].value}
    assert (values["Source rows"], values["Imported rows"], values["Ignored rows"]) == (5, 3, 2)
    assert values["Rows needing attention"] == 0
    assert Decimal(str(values["Opening balance"])) == Decimal("1000.00")
    assert Decimal(str(values["Net movement"])) == Decimal("199.60")
    assert Decimal(str(values["Expected closing"])) == Decimal("1199.60")
    assert Decimal(str(values["Actual closing"])) == Decimal("1187.20")
    assert Decimal(str(values["Difference"])) == Decimal("-12.40")
    assert values["Status"] == "Needs review"

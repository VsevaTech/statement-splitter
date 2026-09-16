"""End-to-end HTTP tests through the FastAPI app (no AI, temp SQLite)."""

from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from sqlalchemy import select

from app import categories as c
from app.db import MerchantRule, PendingUpload, Statement, Transaction
from tests.conftest import DEMO, SIMPLE_CSV, csv_bytes, xlsx_bytes


def _upload(client, name: str, data: bytes):
    resp = client.post("/upload", files={"file": (name, data)}, follow_redirects=False)
    assert resp.status_code == 303, resp.text
    return resp.headers["location"]


def _statement_id(location: str) -> str:
    match = re.fullmatch(r"/statements/([0-9a-f-]{36})", location)
    assert match, location
    return match.group(1)


def _tx_by_desc(session, statement_id: str, needle: str) -> Transaction:
    return session.scalar(
        select(Transaction).where(Transaction.statement_id == statement_id, Transaction.description.contains(needle))
    )


def test_home_shows_privacy_note_and_ai_unavailable(client):
    html = client.get("/").text
    assert "AI suggestions are optional." in html
    assert "Only unknown merchant descriptions are sent to the configured AI provider." in html
    assert "AI suggestions unavailable" in html
    assert "Use AI suggestions" in html


def test_upload_csv_auto_mapped_and_reviewed(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    page = client.get(f"/statements/{sid}")
    assert page.status_code == 200
    assert "GOOGLE*GSUITE 839291" in page.text
    assert "Needs review" in page.text
    txs = session.scalars(select(Transaction).where(Transaction.statement_id == sid)).all()
    by_desc = {t.description: t for t in txs}
    assert by_desc["UBER *TRIP 123456"].category == c.TRANSPORT
    assert by_desc["INCOMING TRANSFER ACME LTD INV-1001"].category == c.INCOME
    assert by_desc["BANK FEE ACCOUNT MAINTENANCE"].category == c.BANK_FEES
    assert by_desc["CERCLI LTD 778812"].category is None
    assert by_desc["MYSTERY SHOP 42"].currency == "EUR"
    # raw rows were discarded once the statement was created
    assert session.scalars(select(PendingUpload)).all() == []


def test_upload_xlsx_with_debit_credit(client, session):
    data = xlsx_bytes(
        ["Transaction Date", "Details", "Debit", "Credit", "Currency"],
        [["2024-09-02", "AWS EMEA 123", 120.5, None, "USD"], ["2024-09-03", "CLIENT PAYMENT", None, 5000, "ILS"]],
    )
    sid = _statement_id(_upload(client, "bank.xlsx", data))
    txs = session.scalars(
        select(Transaction).where(Transaction.statement_id == sid).order_by(Transaction.row_index)
    ).all()
    assert [(t.amount, t.direction, t.category) for t in txs] == [
        (-120.5, "debit", c.CLOUD),
        (5000.0, "credit", c.INCOME),
    ]


def test_column_mapping_ui_when_detection_fails(client, session):
    text = "c1,c2,c3,c4\nfoo,2024-09-01,WOLT IL,-45.90\nbar,2024-09-02,UBER,-30.00\n"
    # First column is junk text, third looks like a description, but headers are meaningless and
    # amount/date detection relies on content. Force low confidence by making dates unparsable.
    text = text.replace("2024-09-01", "yesterday").replace("2024-09-02", "today")
    location = _upload(client, "weird.csv", csv_bytes(text))
    assert location.startswith("/uploads/") and location.endswith("/map")
    form = client.get(location)
    assert form.status_code == 200 and "Map columns" in form.text and "WOLT IL" in form.text  # preview shown

    pending_id = location.split("/")[2]
    resp = client.post(
        location,
        data={"date": "c2", "description": "c3", "amount": "c4", "default_currency": "usd"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    sid = _statement_id(resp.headers["location"])
    txs = session.scalars(select(Transaction).where(Transaction.statement_id == sid)).all()
    assert len(txs) == 2 and {t.currency for t in txs} == {"USD"}
    assert session.get(PendingUpload, pending_id) is None


def test_column_mapping_incomplete_redirects_back(client):
    text = "c1,c2\nyesterday,foo\ntoday,bar\n"
    location = _upload(client, "weird.csv", csv_bytes(text))
    resp = client.post(location, data={"date": "c1"}, follow_redirects=False)
    assert resp.status_code == 303 and "error=" in resp.headers["location"]


def test_filters(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    review = client.get(f"/statements/{sid}/table?filter=review").text
    assert "CERCLI LTD 778812" in review and "MYSTERY SHOP 42" in review and "UBER *TRIP" not in review
    income = client.get(f"/statements/{sid}/table?filter=income").text
    assert "INCOMING TRANSFER" in income and "UBER *TRIP" not in income
    expenses = client.get(f"/statements/{sid}/table?filter=expenses").text
    assert "UBER *TRIP" in expenses and "INCOMING TRANSFER" not in expenses


def test_category_correction_saves_rule_and_updates_summary(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    tx = _tx_by_desc(session, sid, "MYSTERY SHOP")
    resp = client.post(
        f"/transactions/{tx.id}/category", data={"category": c.OFFICE, "remember": "on", "filter": "review"}
    )
    assert resp.status_code == 200
    assert "merchant rule saved" in resp.text
    session.expire_all()
    tx = session.get(Transaction, tx.id)
    assert (tx.category, tx.category_source) == (c.OFFICE, "manual")
    rule = session.scalar(select(MerchantRule).where(MerchantRule.merchant_key == "MYSTERY SHOP"))
    assert rule is not None and rule.category == c.OFFICE


def test_category_correction_without_remember_saves_no_rule(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    tx = _tx_by_desc(session, sid, "MYSTERY SHOP")
    resp = client.post(f"/transactions/{tx.id}/category", data={"category": c.PERSONAL, "remember": "off"})
    assert resp.status_code == 200
    assert session.scalars(select(MerchantRule)).all() == []


def test_invalid_category_rejected(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    tx = _tx_by_desc(session, sid, "MYSTERY SHOP")
    assert client.post(f"/transactions/{tx.id}/category", data={"category": "Yachts"}).status_code == 400


def test_apply_to_all_similar_merchants(client, session):
    text = SIMPLE_CSV + "2024-09-08,MYSTERY SHOP 77,-12.00,EUR\n2024-09-09,MYSTERY SHOP 78,-13.00,EUR\n"
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(text)))
    resp = client.post(
        f"/statements/{sid}/apply-similar",
        data={"merchant_key": "MYSTERY SHOP", "category": c.OFFICE, "filter": "all"},
    )
    assert resp.status_code == 200 and "3 transaction(s)" in resp.text
    rows = session.scalars(select(Transaction).where(Transaction.normalized_merchant == "MYSTERY SHOP")).all()
    assert len(rows) == 3 and all(r.category == c.OFFICE and r.category_source == "manual" for r in rows)
    assert session.scalar(select(MerchantRule).where(MerchantRule.merchant_key == "MYSTERY SHOP")).category == c.OFFICE


def test_critical_scenario_learned_merchant_applies_to_next_statement(client, session):
    """upload → CERCLI unknown → user selects Accounting → rule saved → next upload auto-categorized."""
    first = _statement_id(_upload(client, "september.csv", csv_bytes(SIMPLE_CSV)))
    cercli = _tx_by_desc(session, first, "CERCLI")
    assert cercli.category is None and cercli.category_source == "unknown"

    resp = client.post(f"/transactions/{cercli.id}/category", data={"category": c.ACCOUNTING, "remember": "on"})
    assert resp.status_code == 200
    assert session.scalar(select(MerchantRule).where(MerchantRule.merchant_key == "CERCLI")).category == c.ACCOUNTING

    second_csv = (
        "Date,Description,Amount,Currency\n2024-10-04,CERCLI LTD 990011,-450.00,ILS\n2024-10-05,WOLT IL,-60.00,ILS\n"
    )
    second = _statement_id(_upload(client, "october.csv", csv_bytes(second_csv)))
    cercli2 = _tx_by_desc(session, second, "CERCLI")
    assert cercli2.category == c.ACCOUNTING
    assert cercli2.category_source == "rule"
    assert cercli2.confidence == 1.0
    summary_html = client.get(f"/statements/{second}/table?filter=review").text
    assert "CERCLI" not in summary_html


def test_export_xlsx_endpoint(client):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    resp = client.get(f"/statements/{sid}/export")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "categorized-expenses.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Transactions", "Summary"]
    assert wb["Transactions"].max_row == 7
    summary = [tuple(c.value for c in r) for r in wb["Summary"].iter_rows(min_row=2)]
    currencies = {r[3] for r in summary}
    assert currencies == {"ILS", "USD", "EUR"}


def test_ai_toggle_setting(client):
    assert client.get("/health").json()["ai_enabled"] is False
    client.post("/settings/ai", data={"use_ai": "on"}, follow_redirects=False)
    assert client.get("/health").json()["ai_enabled"] is True
    # Turned on without a key: app keeps working, unknown merchants simply stay in review.
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    assert "CERCLI LTD 778812" in client.get(f"/statements/{sid}/table?filter=review").text
    client.post("/settings/ai", data={}, follow_redirects=False)
    assert client.get("/health").json()["ai_enabled"] is False


def test_delete_statement_and_rules_page(client, session):
    sid = _statement_id(_upload(client, "statement.csv", csv_bytes(SIMPLE_CSV)))
    tx = _tx_by_desc(session, sid, "MYSTERY SHOP")
    client.post(f"/transactions/{tx.id}/category", data={"category": c.OFFICE, "remember": "on"})
    assert "MYSTERY SHOP" in client.get("/rules").text
    assert client.post(f"/statements/{sid}/delete", follow_redirects=False).status_code == 303
    assert session.get(Statement, sid) is None
    assert session.scalars(select(Transaction)).all() == []
    rule = session.scalar(select(MerchantRule))
    assert client.post(f"/rules/{rule.id}/delete", follow_redirects=False).status_code == 303
    assert session.scalars(select(MerchantRule)).all() == []


def test_demo_scenario_end_to_end(client, session):
    """The live demo: statement.xlsx → correct 3 merchants → statement-2.csv has them auto-categorized."""
    sid = _statement_id(_upload(client, "statement.xlsx", (DEMO / "statement.xlsx").read_bytes()))
    txs = session.scalars(select(Transaction).where(Transaction.statement_id == sid)).all()
    assert len(txs) == 300
    auto = sum(1 for t in txs if t.category is not None)
    assert auto / len(txs) > 0.6

    fixes = {"CERCLI": c.ACCOUNTING, "DOCUSIGN": c.SOFTWARE, "TLV PRINTHOUSE": c.OFFICE}
    for key, cat in fixes.items():
        assert any(t.normalized_merchant == key and t.category is None for t in txs), key
        resp = client.post(f"/statements/{sid}/apply-similar", data={"merchant_key": key, "category": cat})
        assert resp.status_code == 200

    assert client.get(f"/statements/{sid}/export").status_code == 200

    sid2 = _statement_id(_upload(client, "statement-2.csv", (DEMO / "statement-2.csv").read_bytes()))
    txs2 = session.scalars(select(Transaction).where(Transaction.statement_id == sid2)).all()
    assert len(txs2) == 140
    for key, cat in fixes.items():
        hits = [t for t in txs2 if t.normalized_merchant == key]
        assert hits, key
        assert all(t.category == cat and t.category_source == "rule" for t in hits), key

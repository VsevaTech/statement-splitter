"""XLSX export: Transactions + Summary (per category AND currency — never mixed)."""

from __future__ import annotations

import io
from collections import defaultdict
from collections.abc import Iterable
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TRANSACTION_HEADERS = (
    "Date",
    "Original Description",
    "Normalized Merchant",
    "Amount",
    "Currency",
    "Direction",
    "Category",
    "Category Source",
)
SUMMARY_HEADERS = ("Category", "Transaction Count", "Total Amount", "Currency")
NEEDS_REVIEW_LABEL = "Needs review"


class TransactionLike(Protocol):
    date: object
    description: str
    normalized_merchant: str
    amount: float
    currency: str
    direction: str
    category: str | None
    category_source: str


def build_summary(transactions: Iterable[TransactionLike]) -> list[tuple[str, int, float, str]]:
    """Rows of (category, count, total, currency), sorted by currency then category.

    Amounts of different currencies are never added together.
    """
    buckets: dict[tuple[str, str], list[float]] = defaultdict(list)
    for tx in transactions:
        buckets[(tx.category or NEEDS_REVIEW_LABEL, tx.currency)].append(tx.amount)
    rows = [(cat, len(v), round(sum(v), 2), cur) for (cat, cur), v in buckets.items()]
    return sorted(rows, key=lambda r: (r[3], r[0]))


def _style_header(ws, width_hint: dict[int, int]) -> None:
    fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for idx, width in width_hint.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def export_workbook(transactions: list[TransactionLike]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(TRANSACTION_HEADERS)
    for tx in transactions:
        date_value = tx.date
        if hasattr(date_value, "isoformat"):
            date_value = date_value.isoformat()
        ws.append(
            [
                date_value or "",
                tx.description,
                tx.normalized_merchant,
                float(tx.amount),
                tx.currency,
                tx.direction,
                tx.category or NEEDS_REVIEW_LABEL,
                tx.category_source,
            ]
        )
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"
    _style_header(ws, {1: 12, 2: 40, 3: 28, 4: 14, 5: 10, 6: 10, 7: 22, 8: 16})
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Summary")
    ws2.append(SUMMARY_HEADERS)
    for category, count, total, currency in build_summary(transactions):
        ws2.append([category, count, total, currency])
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "#,##0.00"
    _style_header(ws2, {1: 24, 2: 18, 3: 16, 4: 10})

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

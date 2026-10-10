"""Read CSV / XLSX statements into a plain table (list of column names + list of rows).

Everything downstream works with plain Python values so that raw rows can be
stored temporarily as JSON while the user confirms the column mapping.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.config import settings

# Only formats we can actually parse. Legacy binary .xls is NOT supported: openpyxl reads the
# zipped OOXML format only, so listing .xls here would accept files that can never be read.
SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".txt")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls / .doc container
_MAX_ZIP_ENTRIES = 2000
LEGACY_XLS_MESSAGE = (
    "Legacy .xls files are not supported. Please re-save the statement as .xlsx or .csv and upload it again."
)
_CANDIDATE_DELIMITERS = [",", ";", "\t", "|"]
_FALLBACK_ENCODINGS = ["utf-8-sig", "utf-8", "cp1255", "cp1251", "cp1252", "latin-1"]
_SINGLE_BYTE_CANDIDATES = ["cp1251", "cp1255", "cp1252", "iso-8859-8", "koi8-r", "cp1250"]


class StatementImportError(ValueError):
    """Raised when a file cannot be parsed into a table."""


@dataclass(slots=True)
class RawTable:
    columns: list[str]
    rows: list[list[str]]
    # Source row number (1-based line in the CSV / row in the sheet) of each entry in `rows`.
    row_numbers: list[int] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None
    sheet: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return len(self.rows)

    def preview(self, n: int = 10) -> list[list[str]]:
        return self.rows[:n]


def _check_table_limits(rows: int, columns: int) -> None:
    if columns > settings.max_columns:
        raise StatementImportError(f"The file has too many columns (limit: {settings.max_columns}).")
    if rows > settings.max_rows:
        raise StatementImportError(f"The file has too many rows (limit: {settings.max_rows}).")


def check_xlsx_container(data: bytes) -> None:
    """Validate the zip container *before* openpyxl inflates it.

    A few KB of zip can expand to gigabytes, so the declared uncompressed sizes are summed
    and compared with the limit. Malformed archives are rejected with a normal import error.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > _MAX_ZIP_ENTRIES:
                raise StatementImportError("The Excel file is not a valid workbook (too many internal files).")
            total = sum(i.file_size for i in infos)
            limit = settings.max_xlsx_uncompressed_mb * 1024 * 1024
            if total > limit:
                raise StatementImportError(
                    f"The Excel file expands to more than {settings.max_xlsx_uncompressed_mb} MB uncompressed, "
                    "which is over the limit."
                )
    except zipfile.BadZipFile as exc:
        raise StatementImportError("The file is not a valid .xlsx workbook (corrupted or not an Excel file).") from exc


def detect_encoding(data: bytes) -> str:
    """Return a usable text encoding for `data`. Conservative: utf-8 first."""
    for enc in ("utf-8-sig", "utf-8"):
        try:
            data.decode(enc)
            return enc
        except UnicodeDecodeError:
            pass
    try:
        from charset_normalizer import from_bytes

        # Restrict to code pages actually seen in bank exports; unrestricted detection
        # is unreliable on small files.
        best = from_bytes(data, cp_isolation=_SINGLE_BYTE_CANDIDATES).best()
        if best is not None and best.encoding:
            return best.encoding
    except Exception:  # pragma: no cover - optional dependency path
        pass
    for enc in _FALLBACK_ENCODINGS[2:]:
        try:
            data.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def detect_delimiter(text: str) -> str:
    """Detect the CSV delimiter from the first lines of `text`."""
    sample_lines = [ln for ln in text.splitlines()[:20] if ln.strip()]
    sample = "\n".join(sample_lines)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(_CANDIDATE_DELIMITERS))
        if dialect.delimiter in _CANDIDATE_DELIMITERS:
            return dialect.delimiter
    except csv.Error:
        pass
    # Fallback: the delimiter that appears consistently (same count) on most lines.
    best, best_score = ",", -1
    for delim in _CANDIDATE_DELIMITERS:
        counts = [ln.count(delim) for ln in sample_lines]
        if not counts or max(counts) == 0:
            continue
        common = max(set(counts), key=counts.count)
        if common == 0:
            continue
        score = counts.count(common) * 1000 + common
        if score > best_score:
            best, best_score = delim, score
    return best


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)  # shortest round-trip text, parsed later as an exact Decimal
    if hasattr(value, "isoformat"):
        try:
            iso = value.isoformat()
            return iso[:10] if iso.endswith("T00:00:00") else iso
        except Exception:  # pragma: no cover
            return str(value)
    return str(value).strip()


def _unique_headers(raw: list[str]) -> list[str]:
    """Column names like pandas used to produce: blank → ``column_N``, duplicates → ``Name.1``."""
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, header in enumerate(raw):
        name = header.strip() if header.strip() and not header.startswith("Unnamed") else f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


def _build_table(
    records: list[tuple[int, list[str]]], *, drop_empty_columns: bool
) -> tuple[list[str], list[list[str]], list[int]]:
    """records: (source row number, cleaned cells). First non-empty record is the header.

    Every record after the header is kept — including blank ones — so each source
    row can be accounted for by the import audit. Only blank rows *after* the
    last non-empty row are dropped (spreadsheet padding, trailing newlines).
    """
    records = [(n, cells) for n, cells in records]
    start = next((i for i, (_n, cells) in enumerate(records) if any(cells)), None)
    if start is None:
        return [], [], []
    header = records[start][1]
    body = records[start + 1 :]
    while body and not any(body[-1][1]):
        body.pop()
    width = len(header)
    # Trailing empty cells beyond the header (e.g. "a,b,c,") are not data.
    body = [(n, cells[:width] if not any(cells[width:]) else cells) for n, cells in body]
    body = [(n, cells + [""] * (width - len(cells))) for n, cells in body]

    keep = list(range(width))
    if drop_empty_columns:
        keep = [i for i in keep if any(i < len(c) and c[i] for _n, c in body)]
    else:
        keep = [i for i in keep if header[i] or any(i < len(c) and c[i] for _n, c in body)]
    extra = lambda cells: cells[width:]  # noqa: E731 - cells beyond the header stay visible
    columns = _unique_headers([header[i] for i in keep])
    rows = [[cells[i] for i in keep] + extra(cells) for _n, cells in body]
    return columns, rows, [n for n, _cells in body]


def read_csv_bytes(data: bytes) -> RawTable:
    encoding = detect_encoding(data)
    text = data.decode(encoding, errors="replace")
    if not text.strip():
        raise StatementImportError("The file is empty.")
    delimiter = detect_delimiter(text)
    records: list[tuple[int, list[str]]] = []
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    line = 1
    try:
        for cells in reader:
            if len(records) >= settings.max_rows + 1:  # +1: the header row
                raise StatementImportError(f"The file has too many rows (limit: {settings.max_rows}).")
            if len(cells) > settings.max_columns:
                raise StatementImportError(f"The file has too many columns (limit: {settings.max_columns}).")
            records.append((line, [_clean_cell(c) for c in cells]))
            line = reader.line_num + 1
    except csv.Error as exc:
        raise StatementImportError(f"Could not parse CSV: {exc}") from exc
    columns, rows, numbers = _build_table(records, drop_empty_columns=False)
    if len(columns) < 2:
        raise StatementImportError("Could not detect a table with at least two columns.")
    return RawTable(columns=columns, rows=rows, row_numbers=numbers, encoding=encoding, delimiter=delimiter)


def read_xlsx_bytes(data: bytes, sheet: str | int = 0) -> RawTable:
    check_xlsx_container(data)
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        # Declared sheet dimensions can be huge for a nearly empty sheet: check before iterating.
        _check_table_limits(ws.max_row or 0, ws.max_column or 0)
        records = [
            (number, [_clean_cell(v) for v in values])
            for number, values in enumerate(ws.iter_rows(values_only=True), start=1)
        ]
    except StatementImportError:
        raise
    except Exception as exc:
        raise StatementImportError("Could not read the Excel file. It may be corrupted or password-protected.") from exc
    columns, rows, numbers = _build_table(records, drop_empty_columns=True)
    if len(columns) < 2 or not rows:
        raise StatementImportError("The first sheet does not contain a table with at least two columns.")
    return RawTable(columns=columns, rows=rows, row_numbers=numbers, sheet=str(sheet))


def read_statement(filename: str, data: bytes) -> RawTable:
    """Parse an uploaded statement. Only ``StatementImportError`` ever escapes."""
    try:
        return _read_statement(filename, data)
    except StatementImportError:
        raise
    except Exception as exc:  # never leak a traceback / 500 for a malformed file
        raise StatementImportError("The file could not be read as a bank statement.") from exc


def _read_statement(filename: str, data: bytes) -> RawTable:
    name = filename.lower()
    if name.endswith(".xls") or data[:8] == _OLE_MAGIC:
        raise StatementImportError(LEGACY_XLS_MESSAGE)
    if name.endswith(".xlsx"):
        return read_xlsx_bytes(data)
    if name.endswith((".csv", ".txt")):
        return read_csv_bytes(data)
    # Unknown extension: sniff by magic bytes (xlsx is a zip container).
    if data[:2] == b"PK":
        return read_xlsx_bytes(data)
    return read_csv_bytes(data)
